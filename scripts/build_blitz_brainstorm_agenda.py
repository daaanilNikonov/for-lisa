#!/usr/bin/env python3
"""Повестка мозгового штурма Блиц-запись — Word, корпоративный стиль Форус."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentation" / "Блиц-запись_Повестка_мозговой_штурм.docx"

BLUE = RGBColor(0x26, 0xA6, 0xE0)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x76, 0x76, 0x77)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run_font(run, size=11, bold=False, color=DARK, name="Verdana"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def shade_cell(cell, hex_color: str):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_cell_borders(cell, color="BFBFBF"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        borders.append(el)
    tc_pr.append(borders)


def add_para(doc_or_cell, text, size=11, bold=False, color=DARK, space_after=6, space_before=0, align=None):
    if hasattr(doc_or_cell, "paragraphs") and hasattr(doc_or_cell, "add_paragraph"):
        # Document or cell
        p = doc_or_cell.add_paragraph()
    else:
        p = doc_or_cell.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_heading_bar(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, "26A6E0")
    set_cell_borders(cell, "26A6E0")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=13, bold=True, color=WHITE)
    doc.add_paragraph()


def decision_box(doc, title="Решения по блоку"):
    add_para(doc, title, size=11, bold=True, color=BLUE, space_before=8, space_after=4)
    table = doc.add_table(rows=4, cols=2)
    table.autofit = True
    headers = ["Что решили", "Ответственный / срок"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        shade_cell(cell, "1A1A1A")
        set_cell_borders(cell)
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=10, bold=True, color=WHITE)
    for r in range(1, 4):
        for c in range(2):
            cell = table.cell(r, c)
            shade_cell(cell, "F5F5F5" if r % 2 else "FFFFFF")
            set_cell_borders(cell)
            p = cell.paragraphs[0]
            run = p.add_run(" " * 2)
            set_run_font(run, size=11, color=GRAY)
            # высота строки через пустые абзацы
            for _ in range(2):
                cell.add_paragraph()
    doc.add_paragraph()


def questions_block(doc, questions: list[str]):
    add_para(doc, "Вопросы для обсуждения", size=11, bold=True, color=BLUE, space_before=4, space_after=4)
    for q in questions:
        add_para(doc, f"•  {q}", size=11, color=DARK, space_after=3)


def build():
    doc = Document()

    # поля страницы
    for section in doc.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    # Титул
    add_para(doc, "ГК Форус  ·  Группа продуктового запуска", size=10, color=GRAY, space_after=2)
    add_para(doc, "Блиц-запись", size=22, bold=True, color=BLUE, space_after=2)
    add_para(doc, "Повестка мозгового штурма", size=16, bold=True, color=DARK, space_after=8)

    # Мета
    meta = doc.add_table(rows=5, cols=2)
    meta_rows = [
        ("Формат", "Очный / онлайн мозговой штурм"),
        ("Длительность", "2 часа"),
        ("Дата", "____________________"),
        ("Участники", "____________________"),
        ("Фасилитатор / секретарь", "____________________"),
    ]
    for i, (k, v) in enumerate(meta_rows):
        ck, cv = meta.cell(i, 0), meta.cell(i, 1)
        shade_cell(ck, "F2F2F2")
        set_cell_borders(ck)
        set_cell_borders(cv)
        pk = ck.paragraphs[0]
        rk = pk.add_run(k)
        set_run_font(rk, size=10, bold=True, color=DARK)
        pv = cv.paragraphs[0]
        rv = pv.add_run(v)
        set_run_font(rv, size=10, color=DARK)
    doc.add_paragraph()

    # Цель
    add_heading_bar(doc, "Цель встречи")
    add_para(
        doc,
        "За 2 часа зафиксировать рабочие решения по четырём направлениям развития Блиц-записи: "
        "конкурентная отстройка, каналы привлечения, сопровождение после продажи и модель демонстрации. "
        "На выходе — список решений с ответственными и сроками, без отложенных «обсудим позже» по критичным пунктам.",
        size=11,
        space_after=10,
    )

    # Тайминг
    add_heading_bar(doc, "Тайминг (ориентир на 120 минут)")
    timing = [
        ("0:00–0:10", "Рамка встречи", "Цель, правила, ожидаемый результат"),
        ("0:10–0:40", "Блок 1", "Отстройка от Dikidi и Masters при бесплатном оффере"),
        ("0:40–1:05", "Блок 2", "Продвижение: digital и событийные контексты"),
        ("1:05–1:25", "Блок 3", "Сопровождение клиента после продажи"),
        ("1:25–1:50", "Блок 4", "Продающая модель демонстрации"),
        ("1:50–2:00", "Сводка", "Решения, владельцы, следующие шаги"),
    ]
    ttable = doc.add_table(rows=1 + len(timing), cols=3)
    for i, h in enumerate(["Время", "Блок", "Содержание"]):
        cell = ttable.cell(0, i)
        shade_cell(cell, "26A6E0")
        set_cell_borders(cell, "26A6E0")
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, size=10, bold=True, color=WHITE)
    for r, (a, b, c) in enumerate(timing, start=1):
        for i, val in enumerate((a, b, c)):
            cell = ttable.cell(r, i)
            shade_cell(cell, "F7F7F7" if r % 2 else "FFFFFF")
            set_cell_borders(cell)
            run = cell.paragraphs[0].add_run(val)
            set_run_font(run, size=10, bold=(i == 1), color=DARK)
    doc.add_paragraph()

    # -------- Блок 1 --------
    add_heading_bar(doc, "Блок 1. Конкурентная отстройка: Dikidi и Masters")
    add_para(doc, "Тезис для обсуждения", size=11, bold=True, color=BLUE, space_after=4)
    add_para(
        doc,
        "Оба конкурента дают бесплатный вход. Значит, наш оффер «бесплатное сотрудничество» сам по себе "
        "не отличает продукт. Нужно определить, чем Блиц-запись выигрывает при равном пороге входа: "
        "в сообщении на первом контакте, в скрипте, в демо и в условиях перехода на платный пакет.",
        size=11,
        space_after=6,
    )
    questions_block(
        doc,
        [
            "Какую одну фразу ценности говорим вместо «у нас тоже бесплатно»?",
            "Где мы реально сильнее Dikidi Free и Masters — и можем ли это доказать на первом звонке?",
            "Что меняем в оффере перехода на платный пакет, чтобы не конкурировать только ценой нуля?",
            "Какие возражения «нам и так бесплатно удобно» закрываем обязательно?",
        ],
    )
    decision_box(doc)

    # -------- Блок 2 --------
    add_heading_bar(doc, "Блок 2. Продвижение: digital и событийные контексты")
    add_para(doc, "Тезис для обсуждения", size=11, bold=True, color=BLUE, space_after=4)
    add_para(
        doc,
        "Нужна рабочая схема привлечения вне холодного прозвона: какие интернет-каналы и какие "
        "офлайн/событийные точки даём в тест в первую очередь, с каким сообщением, бюджетом "
        "и критерием успеха на пилоте.",
        size=11,
        space_after=6,
    )
    questions_block(
        doc,
        [
            "Какие 1–2 digital-канала запускаем в пилот (поиск, карты, соцсети, партнёрские рассылки)?",
            "Какие событийные форматы нам подходят (отраслевые встречи, вебинары, партнёрские дни)?",
            "Какое сообщение и креатив тестируем первым — и кого считаем целевым откликом?",
            "Как измеряем пилот: лиды, демо, стоимость демо, срок до первого сигнала?",
        ],
    )
    decision_box(doc)

    # -------- Блок 3 --------
    add_heading_bar(doc, "Блок 3. Сопровождение клиента после продажи")
    add_para(doc, "Тезис для обсуждения", size=11, bold=True, color=BLUE, space_after=4)
    add_para(
        doc,
        "После сделки клиенту нужна понятная точка коммуникации: где ведём историю, где создаём чаты, "
        "кто отвечает на каком этапе (онбординг, поддержка, развитие). Цель — выбрать основной контур "
        "и правила, чтобы клиент не «терялся» между мессенджерами и CRM.",
        size=11,
        space_after=6,
    )
    questions_block(
        doc,
        [
            "Где храним карточку и статус клиента как систему учёта (CRM / другое)?",
            "Где создаём оперативные чаты с клиентом — и это один канал или несколько по ролям?",
            "Кто владелец чата на онбординге и кто подключается при эскалации?",
            "Какие события обязательно логируем после продажи (старт, обучение, первый успех, риск оттока)?",
        ],
    )
    decision_box(doc)

    # -------- Блок 4 --------
    add_heading_bar(doc, "Блок 4. Модель демонстрации: более продающий формат")
    add_para(doc, "Тезис для обсуждения", size=11, bold=True, color=BLUE, space_after=4)
    add_para(
        doc,
        "Демо должно вести к решению, а не только показывать интерфейс. Нужна единая модель: "
        "структура встречи, акцент на боли клиента, момент оффера и следующий шаг "
        "(тест, подключение, повторный контакт).",
        size=11,
        space_after=6,
    )
    questions_block(
        doc,
        [
            "Какая структура демо на 30–40 минут: вступление, боль, показ, оффер, next step?",
            "Что обязательно выявляем до демо (текущий сервис, штат, поддержка, сценарий записи)?",
            "Какой результат демо считаем успешным — и что делаем, если клиент «подумает»?",
            "Нужны ли разные сценарии демо для Dikidi / Masters / Юклаенс / «пишут вручную»?",
        ],
    )
    decision_box(doc)

    # Сводка
    add_heading_bar(doc, "Итоговая сводка встречи")
    add_para(doc, "Топ-решения дня (заполнить в конце)", size=11, bold=True, color=BLUE, space_after=4)
    summary = doc.add_table(rows=6, cols=3)
    for i, h in enumerate(["№", "Решение", "Владелец / срок"]):
        cell = summary.cell(0, i)
        shade_cell(cell, "1A1A1A")
        set_cell_borders(cell)
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, size=10, bold=True, color=WHITE)
    for r in range(1, 6):
        for c, val in enumerate((str(r), "", "")):
            cell = summary.cell(r, c)
            shade_cell(cell, "FFFFFF")
            set_cell_borders(cell)
            run = cell.paragraphs[0].add_run(val if c == 0 else " ")
            set_run_font(run, size=11, color=DARK)
            if c > 0:
                for _ in range(1):
                    cell.add_paragraph()
    doc.add_paragraph()

    add_para(doc, "Что сознательно не решили / вынесли на следующую встречу", size=11, bold=True, color=BLUE, space_after=4)
    open_t = doc.add_table(rows=4, cols=1)
    for r in range(4):
        cell = open_t.cell(r, 0)
        shade_cell(cell, "F5F5F5")
        set_cell_borders(cell)
        cell.paragraphs[0].add_run(" ")
        for _ in range(1):
            cell.add_paragraph()
    doc.add_paragraph()

    add_para(
        doc,
        "Документ для заполнения на встрече. Стиль оформления: ГК Форус (синий / чёрный / белый).",
        size=9,
        color=GRAY,
        space_after=0,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    build()
