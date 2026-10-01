#!/usr/bin/env python3
"""
Обучение МПП: этап «Сбор потребности» на основе скрипта «Блиц-запись».
Корпоративный стиль — тёмный шаблон ГК Форус 16×9.
"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "Презентация ГК Форус темный шаблон 16х9 (1).pptx"
OUT = ROOT / "presentation" / "Блиц-запись_Сбор_потребности_обучение.pptx"
ICONS = ROOT / "presentation" / "assets" / "icons"

BLUE = RGBColor(0x26, 0xA6, 0xE0)
CARD = RGBColor(0x3F, 0x3F, 0x3F)
CARD_LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
GRAY = RGBColor(0x76, 0x76, 0x77)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xBF, 0xBF, 0xBF)
NEAR_BLACK = RGBColor(0x1A, 0x1A, 0x1A)
TEAL = RGBColor(0x2A, 0x9D, 0x8F)
ORANGE = RGBColor(0xE7, 0x6F, 0x51)
ROW_A = RGBColor(0x2A, 0x2A, 0x2A)
ROW_B = RGBColor(0x1F, 0x1F, 0x1F)
ACCENT_SOFT = RGBColor(0x1E, 0x4A, 0x5C)

FONT = "Verdana"

L_TITLE = 0
L_CONTENT = 3
L_CONTENT2 = 4
L_BG = 6
L_EMPTY = 22


def emu(inches: float) -> int:
    return int(Inches(inches))


def delete_all_slides(prs: Presentation) -> None:
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        r_id = sld_id.get(qn("r:id"))
        prs.part.drop_rel(r_id)
        sld_id_lst.remove(sld_id)


def set_run(run, text, size_pt, bold=False, color=WHITE, font_name=FONT):
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    r_pr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = r_pr.find(qn(tag))
        if el is None:
            el = etree.SubElement(r_pr, qn(tag))
        el.set("typeface", font_name)


def set_anchor(text_frame, anchor=MSO_ANCHOR.TOP):
    body_pr = text_frame._txBody.find(qn("a:bodyPr"))
    if body_pr is not None:
        body_pr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}.get(
                anchor, "t"
            ),
        )


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text,
    size_pt=14,
    bold=False,
    color=WHITE,
    align=PP_ALIGN.LEFT,
    font_name=FONT,
    anchor=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    set_anchor(tf, anchor)
    tf.paragraphs[0].alignment = align
    set_run(tf.paragraphs[0].add_run(), text, size_pt, bold, color, font_name)
    return box


def add_multiline(
    slide,
    left,
    top,
    width,
    height,
    lines,
    size_pt=13,
    bold=False,
    color=WHITE,
    align=PP_ALIGN.LEFT,
    spacing=6,
):
    """lines: list of (text, bold?, color?) or plain strings."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    set_anchor(tf, MSO_ANCHOR.TOP)
    for i, item in enumerate(lines):
        if isinstance(item, str):
            text, is_bold, col = item, bold, color
        else:
            text = item[0]
            is_bold = item[1] if len(item) > 1 else bold
            col = item[2] if len(item) > 2 else color
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if i > 0:
            p.space_before = Pt(spacing)
        set_run(p.add_run(), text, size_pt, is_bold, col)
    return box


def add_card(slide, left, top, width, height, fill=CARD, corner=0.08):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    try:
        shape.adjustments[0] = corner
    except Exception:
        pass
    return shape


def add_icon(slide, name: str, left, top, width, height):
    path = ICONS / name
    if not path.exists():
        return None
    return slide.shapes.add_picture(str(path), left, top, width, height)


def clear_body_placeholders(slide):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx in (13, 14, 1, 2):
            if ph.has_text_frame:
                ph.text_frame.clear()


def fill_title(slide, text, size=26):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 0:
            tf = ph.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            set_run(p.add_run(), text, size, True, WHITE)
            return ph
    return add_textbox(slide, emu(0.97), emu(0.48), emu(11.4), emu(1.0), text, size, True, WHITE)


def add_footer_note(slide, text):
    add_textbox(
        slide,
        emu(0.97),
        emu(6.85),
        emu(11.5),
        emu(0.35),
        text,
        10,
        False,
        GRAY,
    )


# ─── Slides ─────────────────────────────────────────────────────────────────


def slide_title(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_TITLE])
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 0:
            tf = ph.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            set_run(p.add_run(), "Блиц-обучение для МПП", 18, False, SOFT)
            p2 = tf.add_paragraph()
            p2.space_before = Pt(8)
            set_run(p2.add_run(), "Сбор потребности", 32, True, WHITE)
            p3 = tf.add_paragraph()
            p3.space_before = Pt(6)
            set_run(p3.add_run(), "на основе скрипта «Блиц-запись»", 20, True, BLUE)
            p4 = tf.add_paragraph()
            p4.space_before = Pt(14)
            set_run(
                p4.add_run(),
                "Как спрашивать клиента без страха\nи вести звонок к демонстрации",
                14,
                False,
                SOFT,
            )
        elif ph.placeholder_format.idx == 1:
            tf = ph.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            set_run(
                p.add_run(),
                "Источник: скрипт холодного звонка «Блиц-запись»\n"
                "Этап 3 · Вопросы → боль → демо 15–20 минут",
                12,
                False,
                SOFT,
            )
    add_icon(slide, "icon_69.png", emu(10.55), emu(0.5), emu(1.8), emu(1.8))
    return slide


def slide_goals(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_CONTENT])
    fill_title(slide, "Зачем это обучение", 26)
    clear_body_placeholders(slide)

    goals = [
        ("icon_58.png", "Понять этап", "Что такое сбор потребности в скрипте «Блиц-запись» и где он стоит в звонке"),
        ("icon_25.png", "Убрать страх", "Вопросы — не «давление», а способ помочь клиенту и себе"),
        ("icon_27.png", "Освоить карту", "8 вопросов скрипта: что спрашиваем и что фиксируем"),
        ("icon_70.png", "Закрепить", "Кейсы за 2 минуты: какие вопросы задать и к какой боли вести"),
    ]
    left0, top0 = emu(0.97), emu(1.55)
    card_w, card_h = emu(5.7), emu(2.3)
    gap_x, gap_y = emu(0.25), emu(0.2)

    for i, (icon, title, body) in enumerate(goals):
        col, row = i % 2, i // 2
        left = left0 + col * (card_w + gap_x)
        top = top0 + row * (card_h + gap_y)
        add_card(slide, left, top, card_w, card_h)
        add_icon(slide, icon, left + emu(0.25), top + emu(0.65), emu(0.7), emu(0.7))
        add_textbox(slide, left + emu(1.15), top + emu(0.35), card_w - emu(1.4), emu(0.4), title, 16, True, BLUE)
        add_textbox(slide, left + emu(1.15), top + emu(0.9), card_w - emu(1.4), emu(1.15), body, 13, False, WHITE)
    return slide


def slide_call_logic(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_BG])
    fill_title(slide, "Логика звонка по скрипту", 26)
    clear_body_placeholders(slide)

    steps = [
        ("1", "Контакт", "Находим ЛПР,\nберём 2–3 минуты"),
        ("2", "Сервис", "20–25 сек про\n«Блиц-запись»"),
        ("3", "Потребность", "Вопросы → боль\n→ что фиксируем"),
        ("4", "Под боль", "Старт 0 ₽ /\nмодули / тариф"),
        ("5", "Демо", "15–20 минут,\nвыбор из двух слотов"),
    ]
    n = len(steps)
    total_w = 11.5
    gap = 0.18
    card_w = (total_w - gap * (n - 1)) / n
    left0, top = 0.97, 1.9

    for i, (num, title, body) in enumerate(steps):
        left = left0 + i * (card_w + gap)
        fill = BLUE if i == 2 else CARD
        add_card(slide, emu(left), emu(top), emu(card_w), emu(3.6), fill, 0.1)
        add_textbox(
            slide, emu(left + 0.15), emu(top + 0.35), emu(card_w - 0.3), emu(0.55),
            num, 28, True, WHITE, PP_ALIGN.CENTER,
        )
        add_textbox(
            slide, emu(left + 0.1), emu(top + 1.15), emu(card_w - 0.2), emu(0.5),
            title, 15, True, WHITE if i == 2 else BLUE, PP_ALIGN.CENTER,
        )
        add_textbox(
            slide, emu(left + 0.12), emu(top + 1.85), emu(card_w - 0.24), emu(1.4),
            body, 12, False, WHITE, PP_ALIGN.CENTER,
        )
        if i < n - 1:
            add_textbox(
                slide, emu(left + card_w - 0.05), emu(top + 1.5), emu(0.3), emu(0.4),
                "→", 18, True, SOFT, PP_ALIGN.CENTER,
            )

    add_textbox(
        slide, emu(0.97), emu(5.8), emu(11.5), emu(0.7),
        "В звонке не продаём подписку — продаём интерес к демонстрации.\n"
        "Сегодня разбираем только этап 3: сбор потребности.",
        13, False, SOFT,
    )
    return slide


def slide_what_is(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_CONTENT])
    fill_title(slide, "Что такое сбор потребности", 26)
    clear_body_placeholders(slide)

    add_card(slide, emu(0.97), emu(1.55), emu(11.5), emu(1.35), ACCENT_SOFT, 0.08)
    add_textbox(
        slide, emu(1.2), emu(1.75), emu(11.0), emu(1.0),
        "Это этап звонка, где клиент говорит больше нас.\n"
        "Мы выясняем: как сейчас принимают запись, чем пользуются, что болит — и подтверждаем услышанное.",
        15, False, WHITE,
    )

    items = [
        ("Не продажа", "Не читаем прайс и не спорим с «Юклиентс» / «Дикиди»"),
        ("Диагностика", "Открытый переход: «Расскажите, как у вас сейчас записываются клиенты?»"),
        ("Фиксация", "Канал · кто ведёт · инструмент · потери · неявки · главная боль"),
        ("Мост к офферу", "После 3.8 («Правильно понимаю…») — рассказ под боль"),
    ]
    for i, (t, b) in enumerate(items):
        left = emu(0.97 + (i % 4) * 2.95)
        top = emu(3.2)
        add_card(slide, left, top, emu(2.8), emu(2.9))
        add_textbox(slide, left + emu(0.15), top + emu(0.3), emu(2.5), emu(0.7), t, 14, True, BLUE, PP_ALIGN.CENTER)
        add_textbox(slide, left + emu(0.15), top + emu(1.15), emu(2.5), emu(1.5), b, 12, False, WHITE, PP_ALIGN.CENTER)
    return slide


def slide_why(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_CONTENT])
    fill_title(slide, "Зачем нужен этот этап", 26)
    clear_body_placeholders(slide)

    reasons = [
        ("Без боли — мимо", "Презентация «в воздух» не цепляет. Боль из ответов клиента делает оффер личным."),
        ("Старт 0 ₽ — не всегда ответ", "Кому-то нужна цена, кому-то — онлайн с 2ГИС, кому-то — порядок в расписании."),
        ("ЛПР и оператор", "Вопрос «кто ведёт расписание» показывает, с кем решать и кого звать на демо."),
        ("Возражения мягче", "Когда клиент сам назвал боль, легче пригласить на 15 минут демо."),
        ("Не спорим с конкурентом", "Спрашиваем «что удобно / что мешает» — сравниваем, а не критикуем."),
        ("Цель звонка — демо", "Сбор потребности готовит приглашение, а не «закрытие» на оплату."),
    ]
    for i, (t, b) in enumerate(reasons):
        col, row = i % 3, i // 3
        left = emu(0.97 + col * 3.9)
        top = emu(1.55 + row * 2.5)
        add_card(slide, left, top, emu(3.7), emu(2.3))
        add_textbox(slide, left + emu(0.2), top + emu(0.3), emu(3.3), emu(0.55), t, 14, True, BLUE)
        add_textbox(slide, left + emu(0.2), top + emu(0.95), emu(3.3), emu(1.15), b, 12, False, WHITE)
    return slide


def slide_no_fear(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_BG])
    fill_title(slide, "Спрашивать без страха", 26)
    clear_body_placeholders(slide)

    myths = [
        ("Миф", "«Я навязываюсь вопросами»", "Факт", "Клиент чувствует интерес.\nМолчание и монолог — хуже."),
        ("Миф", "«Сейчас откажут»", "Факт", "Отказ чаще от продажи\nв лоб, а не от вопроса."),
        ("Миф", "«Надо сразу про тариф»", "Факт", "Сначала боль — потом\nСтарт 0 ₽ под эту боль."),
        ("Миф", "«Не знаю, что спросить»", "Факт", "В скрипте уже есть карта:\n3.1 → 3.8. Не импровизируй с нуля."),
    ]
    for i, (a, at, b, bt) in enumerate(myths):
        col, row = i % 2, i // 2
        left = emu(0.97 + col * 6.0)
        top = emu(1.5 + row * 2.5)
        add_card(slide, left, top, emu(5.75), emu(2.3))
        add_textbox(slide, left + emu(0.25), top + emu(0.25), emu(2.4), emu(0.35), a, 11, True, ORANGE)
        add_textbox(slide, left + emu(0.25), top + emu(0.65), emu(2.4), emu(1.3), at, 13, False, WHITE)
        add_textbox(slide, left + emu(3.0), top + emu(0.25), emu(2.5), emu(0.35), b, 11, True, TEAL)
        add_textbox(slide, left + emu(3.0), top + emu(0.65), emu(2.5), emu(1.3), bt, 13, False, WHITE)

    return slide


def slide_how_ask(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_CONTENT])
    fill_title(slide, "Как задавать вопросы", 26)
    clear_body_placeholders(slide)

    add_card(slide, emu(0.97), emu(1.5), emu(5.7), emu(4.9))
    add_textbox(slide, emu(1.2), emu(1.7), emu(5.2), emu(0.4), "Типы из скрипта", 16, True, BLUE)
    add_multiline(
        slide, emu(1.2), emu(2.3), emu(5.2), emu(3.8),
        [
            ("Открытый", True, BLUE),
            ("Клиент рассказывает. Пример: «Как записываются клиенты?»", False, WHITE),
            ("", False, WHITE),
            ("Закрытый → открытый", True, BLUE),
            ("Сначала факт, потом глубина: «Чем пользуетесь — тетрадь, Юклиентс…?» → «Что мешает?»", False, WHITE),
            ("", False, WHITE),
            ("Подтверждение", True, BLUE),
            ("3.8: «Правильно понимаю: сейчас… мешает… важно… Верно?»", False, WHITE),
        ],
        13,
        spacing=4,
    )

    add_card(slide, emu(6.9), emu(1.5), emu(5.55), emu(4.9))
    add_textbox(slide, emu(7.15), emu(1.7), emu(5.1), emu(0.4), "Правила на звонке", 16, True, BLUE)
    rules = [
        "1. Один вопрос — пауза — слушаем",
        "2. Не дольше 40–50 сек подряд говорим сами",
        "3. Не спорим с текущим сервисом",
        "4. Фиксируем ответы (даже коротко)",
        "5. Не прыгаем сразу в прайс",
        "6. Завершаем повтором услышанного",
        "7. Из боли берём 2–3 пункта в оффер",
    ]
    add_multiline(
        slide, emu(7.15), emu(2.35), emu(5.1), emu(3.7),
        [(r, False, WHITE) for r in rules],
        13,
        spacing=8,
    )
    return slide


def slide_questions_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_CONTENT2])
    fill_title(slide, "Карта вопросов · 3.1–3.4", 26)
    clear_body_placeholders(slide)

    rows = [
        ("3.1", "Как записывают", "Клиенты чаще звонят, пишут в мессенджеры, из соцсетей — или уже есть онлайн-запись?", "Канал записи"),
        ("3.2", "Кто ведёт", "Кто этим занимается — вы сами, администратор или мастер?", "ЛПР / оператор"),
        ("3.3", "Чем пользуются", "Тетрадь, таблица, «Юклиентс», «Дикиди»? → что удобно / что мешает?", "Инструмент + боли"),
        ("3.4", "Потери заявок", "В час пик не дозваниваются, пишут — и запись пропадает? Как часто?", "Потери и хаос"),
    ]
    headers = ["№", "Тема", "Вопрос (из скрипта)", "Фиксируем"]
    widths = [0.7, 2.0, 6.3, 2.3]
    left0, top0 = 0.97, 1.45
    row_h = 1.15

    x = left0
    for h, w in zip(headers, widths):
        add_card(slide, emu(x), emu(top0), emu(w - 0.08), emu(0.45), BLUE, 0.05)
        add_textbox(slide, emu(x + 0.08), emu(top0 + 0.08), emu(w - 0.2), emu(0.3), h, 11, True, WHITE)
        x += w

    for i, (num, theme, q, fix) in enumerate(rows):
        top = top0 + 0.55 + i * row_h
        fill = ROW_A if i % 2 == 0 else ROW_B
        x = left0
        vals = [num, theme, q, fix]
        for v, w in zip(vals, widths):
            add_card(slide, emu(x), emu(top), emu(w - 0.08), emu(row_h - 0.08), fill, 0.04)
            add_textbox(
                slide, emu(x + 0.08), emu(top + 0.15), emu(w - 0.2), emu(row_h - 0.35),
                v, 11 if w > 2 else 12, True if w < 2.2 else False, BLUE if w < 2.2 else WHITE,
            )
            x += w
    add_footer_note(slide, "Источник: скрипт холодного звонка «Блиц-запись», блок 3. Сбор потребности")
    return slide


def slide_questions_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_CONTENT2])
    fill_title(slide, "Карта вопросов · 3.5–3.8", 26)
    clear_body_placeholders(slide)

    rows = [
        ("3.5", "Онлайн и каналы", "Могут сами записаться с сайта, 2ГИС, соцсетей — или всё через администратора? Было бы удобно без звонка?", "Готовность к онлайн"),
        ("3.6", "Неявки", "Вы напоминаете о визите? → А неявки сильно бьют по загрузке?", "Напоминания / no-show"),
        ("3.7", "Главная боль", "Если честно: что в записи и расписании больше всего отнимает время или нервы?", "Главная боль"),
        ("3.8", "Повтор", "Правильно понимаю: сейчас [как], мешает [боль], важно [цель]. Верно?", "Подтверждение"),
    ]
    headers = ["№", "Тема", "Вопрос (из скрипта)", "Фиксируем"]
    widths = [0.7, 2.0, 6.3, 2.3]
    left0, top0 = 0.97, 1.45
    row_h = 1.15

    x = left0
    for h, w in zip(headers, widths):
        add_card(slide, emu(x), emu(top0), emu(w - 0.08), emu(0.45), TEAL, 0.05)
        add_textbox(slide, emu(x + 0.08), emu(top0 + 0.08), emu(w - 0.2), emu(0.3), h, 11, True, WHITE)
        x += w

    for i, (num, theme, q, fix) in enumerate(rows):
        top = top0 + 0.55 + i * row_h
        fill = ROW_A if i % 2 == 0 else ROW_B
        x = left0
        vals = [num, theme, q, fix]
        for v, w in zip(vals, widths):
            add_card(slide, emu(x), emu(top), emu(w - 0.08), emu(row_h - 0.08), fill, 0.04)
            add_textbox(
                slide, emu(x + 0.08), emu(top + 0.15), emu(w - 0.2), emu(row_h - 0.35),
                v, 11 if w > 2 else 12, True if w < 2.2 else False, TEAL if w < 2.2 else WHITE,
            )
            x += w
    add_footer_note(slide, "3.8 — обязательный мост: без подтверждения не переходим к рассказу под боль")
    return slide


def slide_pain_to_pitch(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_CONTENT])
    fill_title(slide, "Из боли — в рассказ (этап 4)", 26)
    clear_body_placeholders(slide)

    pairs = [
        ("Теряют заявки / много звонков", "Сами записываются с сайта, 2ГИС, соцсетей → сразу в календарь"),
        ("Дорого / не готовы платить", "Тариф «Старт» — 0 ₽/мес: календарь, онлайн-запись, прайс, до 5 сотрудников"),
        ("Путаница в расписании", "Одно расписание, без задвоений; на «Управлении» — графики и загрузка"),
        ("Нужна база / возврат", "База на «Старте»; история, аналитика, лояльность — в модулях / «Управление»"),
        ("Всё под ключ и бренд", "«Премиум»: менеджер, брендирование, запуск, каналы и аналитика"),
    ]
    for i, (pain, pitch) in enumerate(pairs):
        top = emu(1.45 + i * 0.95)
        add_card(slide, emu(0.97), top, emu(5.4), emu(0.85))
        add_textbox(slide, emu(1.15), top + emu(0.22), emu(5.05), emu(0.5), pain, 13, True, ORANGE)
        add_textbox(slide, emu(6.5), top + emu(0.25), emu(0.4), emu(0.4), "→", 16, True, SOFT, PP_ALIGN.CENTER)
        add_card(slide, emu(7.0), top, emu(5.45), emu(0.85), ACCENT_SOFT)
        add_textbox(slide, emu(7.2), top + emu(0.15), emu(5.05), emu(0.6), pitch, 12, False, WHITE)
    return slide


def slide_exercise_intro(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_TITLE])
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 0:
            tf = ph.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            set_run(p.add_run(), "Практика", 18, False, SOFT)
            p2 = tf.add_paragraph()
            p2.space_before = Pt(10)
            set_run(p2.add_run(), "Упражнения на закрепление", 30, True, WHITE)
            p3 = tf.add_paragraph()
            p3.space_before = Pt(12)
            set_run(
                p3.add_run(),
                "Таймер · вопросы из скрипта · без чтения прайса\n"
                "Разбор — на следующих слайдах",
                15,
                False,
                BLUE,
            )
        elif ph.placeholder_format.idx == 1:
            tf = ph.text_frame
            tf.clear()
            set_run(
                tf.paragraphs[0].add_run(),
                "Кейсы опираются на реальные сценарии клиентов «Блиц-запись»",
                12,
                False,
                SOFT,
            )
    add_icon(slide, "icon_54.png", emu(10.55), emu(0.5), emu(1.8), emu(1.8))
    return slide


def slide_case_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_BG])
    fill_title(slide, "Упражнение 1 · Кейс за 2 минуты", 24)
    clear_body_placeholders(slide)

    add_card(slide, emu(0.97), emu(1.4), emu(7.4), emu(5.1))
    add_textbox(slide, emu(1.2), emu(1.55), emu(7.0), emu(0.4), "Салон красоты «Луна»", 16, True, BLUE)
    add_multiline(
        slide, emu(1.2), emu(2.1), emu(7.0), emu(4.1),
        [
            "Клиент на линии — владелица салона.",
            "",
            "Факты, которые она уже сказала (или видно из разговора):",
            "• Сейчас сидят на «Юклиентс», 6 мастеров",
            "• База ~2 000 клиентов, ~100 услуг",
            "• «Платим много, а половиной функций не пользуемся»",
            "• Сайт грузится медленно, раздражает",
            "• Спрашивает: «А переносить базу не замучаемся?»",
            "",
            "Задача (2 минуты): выпиши 4–5 вопросов из карты\n"
            "сбора потребности, которые нужно задать дальше,\n"
            "и одну фразу повтора (3.8).",
        ],
        13,
        spacing=3,
    )

    add_card(slide, emu(8.6), emu(1.4), emu(3.85), emu(5.1), BLUE, 0.1)
    add_textbox(slide, emu(8.85), emu(1.7), emu(3.35), emu(0.4), "Формат", 14, True, WHITE)
    add_multiline(
        slide, emu(8.85), emu(2.3), emu(3.35), emu(3.8),
        [
            "1. Таймер 2:00",
            "2. Пиши вопросы, не оффер",
            "3. Можно смотреть карту 3.1–3.8",
            "4. Потом — разбор",
            "",
            "Критерий успеха:",
            "есть вопросы про цену/лишнее,",
            "миграцию, кто ведёт запись,",
            "и подтверждение боли",
        ],
        13,
        spacing=6,
    )
    return slide


def slide_case_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_BG])
    fill_title(slide, "Упражнение 2 · Кейс за 2 минуты", 24)
    clear_body_placeholders(slide)

    add_card(slide, emu(0.97), emu(1.4), emu(7.4), emu(5.1))
    add_textbox(slide, emu(1.2), emu(1.55), emu(7.0), emu(0.4), "Автомойка «Элис» · 8 постов", 16, True, BLUE)
    add_multiline(
        slide, emu(1.2), emu(2.1), emu(7.0), emu(4.1),
        [
            "Клиент — руководитель комплекса.",
            "",
            "Что известно:",
            "• ~2 000 записей в месяц, онлайн — только ~5%",
            "• 95% идёт через администратора (телефон / на месте)",
            "• Недовольны текущим сервисом: ошибки в зарплатах,",
            "  поддержка не помогла с телефонией",
            "• В час пик администраторы «тонут» в потоке",
            "• Цена текущего решения ~120 000 ₽/год",
            "",
            "Задача (2 минуты): какие вопросы задать,\n"
            "чтобы уточнить потребность и не уйти сразу\n"
            "в «у нас дешевле»? Назови главную боль для 3.7.",
        ],
        13,
        spacing=3,
    )

    add_card(slide, emu(8.6), emu(1.4), emu(3.85), emu(5.1), ORANGE, 0.1)
    add_textbox(slide, emu(8.85), emu(1.7), emu(3.35), emu(0.4), "Подсказка", 14, True, WHITE)
    add_multiline(
        slide, emu(8.85), emu(2.3), emu(3.35), emu(3.8),
        [
            "Не начинай с цены.",
            "",
            "Сначала: канал записи,",
            "кто ведёт, час пик,",
            "что отнимает нервы,",
            "нужна ли онлайн-запись",
            "вообще.",
            "",
            "Цена — следствие,",
            "не первый вопрос.",
        ],
        13,
        spacing=5,
    )
    return slide


def slide_case_3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_CONTENT])
    fill_title(slide, "Упражнение 3 · Перепиши вопрос", 24)
    clear_body_placeholders(slide)

    add_textbox(
        slide, emu(0.97), emu(1.4), emu(11.5), emu(0.5),
        "За 90 секунд: замени «плохие» формулировки на вопросы в духе скрипта.",
        14, False, SOFT,
    )

    bad_good = [
        ("«Вам нужна онлайн-запись?»", "→ Открытый / закрытый→открытый из карты 3.1 или 3.5"),
        ("«Сколько вы платите за Юклиентс?»", "→ Сначала 3.3: чем пользуетесь → что мешает / лишнее"),
        ("«Давайте я расскажу про наши тарифы»", "→ Сначала 3.7 и 3.8, потом рассказ под боль"),
        ("«Почему у вас до сих пор тетрадь?»", "→ Без оценки: «Как сейчас принимаете запись?»"),
    ]
    for i, (bad, hint) in enumerate(bad_good):
        top = emu(2.0 + i * 1.1)
        add_card(slide, emu(0.97), top, emu(5.5), emu(0.95))
        add_textbox(slide, emu(1.15), top + emu(0.28), emu(5.15), emu(0.5), bad, 13, True, ORANGE)
        add_card(slide, emu(6.7), top, emu(5.75), emu(0.95), ACCENT_SOFT)
        add_textbox(slide, emu(6.9), top + emu(0.28), emu(5.35), emu(0.5), hint, 12, False, WHITE)
    return slide


def slide_answers_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_CONTENT])
    fill_title(slide, "Разбор · кейс «Луна»", 26)
    clear_body_placeholders(slide)

    add_card(slide, emu(0.97), emu(1.45), emu(11.5), emu(5.1))
    add_multiline(
        slide, emu(1.25), emu(1.65), emu(11.0), emu(4.7),
        [
            ("Сильные вопросы (опора на скрипт):", True, BLUE),
            ("• 3.2 Кто ведёт запись и расписание мастеров — вы или администратор?", False, WHITE),
            ("• 3.3 Что в «Юклиентс» удобно, а что лишнее / мешает? (уже намекнули на цену и функции)", False, WHITE),
            ("• 3.5 Клиенты могут записаться сами с сайта / 2ГИС / соцсетей?", False, WHITE),
            ("• Уточнение боли миграции: «Базу и услуги кто-то уже пробовал переносить — или это главный страх?»", False, WHITE),
            ("• 3.7 Что сильнее бесит: цена, лишние функции или скорость работы?", False, WHITE),
            ("", False, WHITE),
            ("Фраза 3.8:", True, BLUE),
            ("«Правильно понимаю: сейчас «Юклиентс», 6 мастеров, платите за функции, которыми почти не пользуетесь,", False, WHITE),
            ("сайт тормозит, и важно перейти без потери базы ~2 000 клиентов. Верно?»", False, WHITE),
            ("", False, WHITE),
            ("Дальше в оффер: Старт/экономия + переход под ключ / Премиум при страхе миграции — не прайс целиком.", False, TEAL),
        ],
        13,
        spacing=4,
    )
    return slide


def slide_answers_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_CONTENT])
    fill_title(slide, "Разбор · кейс «Элис»", 26)
    clear_body_placeholders(slide)

    add_card(slide, emu(0.97), emu(1.45), emu(11.5), emu(5.1))
    add_multiline(
        slide, emu(1.25), emu(1.65), emu(11.0), emu(4.7),
        [
            ("Сильные вопросы:", True, BLUE),
            ("• 3.1 Как сейчас принимают запись в час пик — звонок, на месте, мессенджер?", False, WHITE),
            ("• 3.2 Кто ведёт 8 постов: сколько администраторов, кто решает по ПО?", False, WHITE),
            ("• 3.4 Бывает, что поток «ломает» запись и заявки теряются? Как часто?", False, WHITE),
            ("• 3.5 Онлайн вам принципиально нужен, или важнее удобный инструмент администратора?", False, WHITE),
            ("• 3.3 / боль: «Зарплаты и поддержка — что больнее прямо сейчас?»", False, WHITE),
            ("• 3.7 Главная боль для формулировки: хаос в пике + ошибки расчётов + глухая поддержка", False, WHITE),
            ("", False, WHITE),
            ("3.8:", True, BLUE),
            ("«Правильно понимаю: 95% записи через администратора, в пике не хватает удобного инструмента,", False, WHITE),
            ("бесят ошибки в зарплатах и поддержка, а онлайн — не главная цель. Важно, чтобы система слушала вас. Верно?»", False, WHITE),
            ("", False, WHITE),
            ("Не вести с «у нас 0 ₽» первым — здесь боль в кастомизации и работе администратора.", False, TEAL),
        ],
        13,
        spacing=4,
    )
    return slide


def slide_cheat(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_CONTENT])
    fill_title(slide, "Шпаргалка на звонок", 26)
    clear_body_placeholders(slide)

    blocks = [
        ("Перед вопросами", "Коротко про сервис (20–25 сек)\n→ «Расскажите, как сейчас записываются клиенты?»"),
        ("Минимум 5 тем", "Канал · кто · инструмент · потери/пик · главная боль"),
        ("Обязательно", "3.8 повтор услышанного\nперед любым оффером"),
        ("После боли", "2–3 пункта из ответов\n→ приглашение на демо 15–20 мин"),
    ]
    for i, (t, b) in enumerate(blocks):
        left = emu(0.97 + (i % 4) * 2.95)
        add_card(slide, left, emu(1.55), emu(2.8), emu(3.5))
        add_textbox(slide, left + emu(0.15), emu(1.75), emu(2.5), emu(0.7), t, 14, True, BLUE, PP_ALIGN.CENTER)
        add_textbox(slide, left + emu(0.15), emu(2.6), emu(2.5), emu(2.1), b, 12, False, WHITE, PP_ALIGN.CENTER)

    add_card(slide, emu(0.97), emu(5.25), emu(11.5), emu(1.2), BLUE, 0.08)
    add_textbox(
        slide, emu(1.2), emu(5.5), emu(11.0), emu(0.7),
        "Помни: спрашивать — не страшно. Страшно продавать вслепую.\n"
        "Скрипт уже дал тебе вопросы — твоя работа — слушать и повторить боль словами клиента.",
        14, True, WHITE, PP_ALIGN.CENTER,
    )
    return slide


def slide_homework(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_CONTENT])
    fill_title(slide, "Итог и практика после обучения", 26)
    clear_body_placeholders(slide)

    items = [
        ("Сегодня", "Знаешь место этапа 3, карту 3.1–3.8 и зачем повтор 3.8"),
        ("На смене", "В ближайших 5 звонках по «Блиц-записи» явно пройди блок вопросов до прайса"),
        ("Самопроверка", "После звонка: какие 3 ответа зафиксировал? Смог ли сказать 3.8 своими словами?"),
        ("С наставником", "Разбор 1 записи звонка: где спросил / где ушёл в монолог"),
    ]
    for i, (t, b) in enumerate(items):
        top = emu(1.5 + i * 1.2)
        add_card(slide, emu(0.97), top, emu(2.4), emu(1.05), BLUE if i == 0 else CARD)
        add_textbox(slide, emu(1.1), top + emu(0.3), emu(2.1), emu(0.5), t, 14, True, WHITE, PP_ALIGN.CENTER)
        add_card(slide, emu(3.55), top, emu(8.9), emu(1.05))
        add_textbox(slide, emu(3.8), top + emu(0.3), emu(8.4), emu(0.55), b, 14, False, WHITE)
    return slide


def slide_close(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[L_TITLE])
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 0:
            tf = ph.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            set_run(p.add_run(), "Вопросы без страха", 18, False, SOFT)
            p2 = tf.add_paragraph()
            p2.space_before = Pt(10)
            set_run(p2.add_run(), "Сначала потребность —\nпотом демонстрация", 28, True, WHITE)
            p3 = tf.add_paragraph()
            p3.space_before = Pt(16)
            set_run(p3.add_run(), "«Блиц-запись» · скрипт холодного звонка · этап 3", 14, False, BLUE)
        elif ph.placeholder_format.idx == 1:
            tf = ph.text_frame
            tf.clear()
            set_run(
                tf.paragraphs[0].add_run(),
                "Материалы: скрипт · блок-схема · кейсы «Луна» и «Элис»",
                12,
                False,
                SOFT,
            )
    add_icon(slide, "icon_69.png", emu(10.55), emu(0.5), emu(1.8), emu(1.8))
    return slide


def main():
    if not TEMPLATE.exists():
        raise SystemExit(f"Template not found: {TEMPLATE}")

    prs = Presentation(str(TEMPLATE))
    delete_all_slides(prs)

    builders = [
        slide_title,
        slide_goals,
        slide_call_logic,
        slide_what_is,
        slide_why,
        slide_no_fear,
        slide_how_ask,
        slide_questions_1,
        slide_questions_2,
        slide_pain_to_pitch,
        slide_exercise_intro,
        slide_case_1,
        slide_case_2,
        slide_case_3,
        slide_answers_1,
        slide_answers_2,
        slide_cheat,
        slide_homework,
        slide_close,
    ]
    for build in builders:
        build(prs)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"Saved: {OUT} ({len(builders)} slides)")


if __name__ == "__main__":
    main()
